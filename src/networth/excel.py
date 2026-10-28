"""Excel export and import in the layout of the original NetWorth.xlsx.

Export: sheet "Net Worth" (date, configured columns, live formulas with cached values, so pandas and the old
script can read it without opening it in Excel first) and sheet "Income" (monthly net salary, yearly sums).

Import: reads such a workbook. Input cells keep typed expressions ("=6604.12+895.2"); formula cells whose value
differs from the configured formula (e.g. older rows with a different formula) are kept as fixed values.
"""
from __future__ import annotations

import datetime as dt
import io
import math
from dataclasses import dataclass, field
from typing import BinaryIO

import openpyxl
import xlsxwriter
from xlsxwriter.utility import xl_col_to_name

from . import formula as fx
from .columns import DATE_COLUMN, Column, parse_highlight
from .db import Cell, Entry, Income
from .sheet import Calculator, Row

SHEET = "Net Worth"
INCOME_SHEET = "Income"
INCOME_COLUMNS = ["Datum", "Nettogehalt", "Jahressumme"]

CURRENCY_FORMAT = '"€"\\ #,##0.00'
PERCENT_FORMAT = "0.00%"
NUMBER_FORMAT = "#,##0.00"
DATE_FORMAT = 14  # built-in short date, shown in the locale of the viewer (e.g. 31.12.2024)
HEADER_FILL = "#FFF2CC"
HIGHLIGHT_FORMATS = {"red": {"bg_color": "#FFC7CE", "font_color": "#9C0006"},
                     "yellow": {"bg_color": "#FFEB9C", "font_color": "#9C5700"},
                     "green": {"bg_color": "#C6EFCE", "font_color": "#006100"}}
EXCEL_OPS = {">": ">", "<": "<", ">=": ">=", "<=": "<=", "=": "==", "<>": "!="}


# ---------------------------------------------------------------------------
# Export
# ---------------------------------------------------------------------------

def export_workbook(columns: list[Column], entries: list[Entry], rows: list[Row], income: list[Income]) -> bytes:
    columns = sorted(columns, key=lambda c: c.position)
    letters = {c.name: xl_col_to_name(i) for i, c in enumerate(columns, start=1)}
    buffer = io.BytesIO()
    wb = xlsxwriter.Workbook(buffer, {"in_memory": True})
    fmt_header = wb.add_format({"bold": True, "bg_color": HEADER_FILL, "bottom": 2})
    fmt_date = wb.add_format({"num_format": DATE_FORMAT})
    fmt_number = {"currency": wb.add_format({"num_format": CURRENCY_FORMAT}),
                  "percent": wb.add_format({"num_format": PERCENT_FORMAT}),
                  "number": wb.add_format({"num_format": NUMBER_FORMAT})}

    ws = wb.add_worksheet(SHEET)
    ws.freeze_panes(1, 1)
    ws.write(0, 0, DATE_COLUMN, fmt_header)
    ws.set_column(0, 0, 11.3)
    for i, col in enumerate(columns, start=1):
        ws.write(0, i, col.name, fmt_header)
        ws.set_column(i, i, col.width)

    parsed = {c.name: c.parsed for c in columns if c.kind == "formula"}
    for r, (entry, row) in enumerate(zip(entries, rows), start=1):
        excel_row = r + 1
        ws.write_datetime(r, 0, dt.datetime.combine(row.date, dt.time()), fmt_date)
        for i, col in enumerate(columns, start=1):
            value = row.values.get(col.name)
            if col.kind == "text":
                if value:
                    ws.write_string(r, i, str(value))
                continue
            fmt = fmt_number[col.number_format]
            if col.kind == "input":
                cell = entry.cells.get(col.id)
                if cell is not None and cell.expr and value is not None:
                    ws.write_formula(r, i, "=" + cell.expr, fmt, value)
                elif value is not None:
                    ws.write_number(r, i, value, fmt)
                else:
                    ws.write_blank(r, i, None, fmt)
            elif col.name in row.overridden or (r == 1 and fx.uses_previous(parsed[col.name])):
                if value is None:
                    ws.write_blank(r, i, None, fmt)
                else:
                    ws.write_number(r, i, value, fmt)
            else:
                # cached result, so the file can be read without Excel; an empty result was an error like 1/0
                cached = value if value is not None else "#DIV/0!"
                ws.write_formula(r, i, "=" + excel_formula(parsed[col.name], letters, excel_row), fmt, cached)

    last_row = max(len(rows), 1)
    for i, col in enumerate(columns, start=1):
        if col.highlight:
            for rule in parse_highlight(col.highlight):
                ws.conditional_format(1, i, max(last_row, 1000), i, {
                    "type": "cell", "criteria": EXCEL_OPS[rule.op], "value": rule.value,
                    "format": wb.add_format(HIGHLIGHT_FORMATS[rule.style])})

    total = next((c for c in columns if c.role == "total"), None)
    wb.define_name("Datumsachse", f"=OFFSET('{SHEET}'!$A$1,0,0,COUNTA('{SHEET}'!$A:$A),1)")
    if total is not None:
        letter = letters[total.name]
        wb.define_name("Summenachse", f"=OFFSET('{SHEET}'!${letter}$1,0,0,COUNTA('{SHEET}'!${letter}:${letter}),1)")

    _write_income(wb, income, fmt_header, fmt_date, fmt_number["currency"])
    wb.close()
    return buffer.getvalue()


def excel_formula(node: fx.Node, letters: dict[str, str], row: int) -> str:
    def ref(r: fx.Ref) -> str:
        letter = letters[r.name]
        if r.mode == "prev":
            return f"{letter}{row - 1}"
        if r.mode == "first":
            return f"${letter}$2"
        return f"{letter}{row}"
    return fx.render(node, ref)


def _write_income(wb, income: list[Income], fmt_header, fmt_date, fmt_currency) -> None:
    ws = wb.add_worksheet(INCOME_SHEET)
    for i, (name, width) in enumerate(zip(INCOME_COLUMNS, (23.6, 11.4, 13.1))):
        ws.write(0, i, name, fmt_header)
        ws.set_column(i, i, width)
    year_start = 1
    for r, item in enumerate(income, start=1):
        ws.write_datetime(r, 0, dt.datetime.combine(item.date, dt.time()), fmt_date)
        ws.write_number(r, 1, item.amount, fmt_currency)
        is_last_of_year = r == len(income) or income[r].date.year != item.date.year
        if is_last_of_year:  # yearly sum next to the last month of every year
            total = math.fsum(i.amount for i in income[year_start - 1:r])
            ws.write_formula(r, 2, f"=SUM(B{year_start + 1}:B{r + 1})", fmt_currency, total)
            year_start = r + 1


# ---------------------------------------------------------------------------
# Import
# ---------------------------------------------------------------------------

@dataclass
class ImportResult:
    entries: list[Entry]
    income: list[Income] | None
    new_columns: list[Column]
    ignored: list[str] = field(default_factory=list)
    fixed_values: int = 0  # formula cells kept as fixed values


def _to_date(value) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str) and value.strip():
        for parse in (dt.date.fromisoformat, lambda s: dt.datetime.strptime(s, "%d.%m.%Y").date()):
            try:
                return parse(value.strip())
            except ValueError:
                pass
    return None


def _to_number(value) -> float | None:
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        return float(value) if math.isfinite(value) else None
    if isinstance(value, str):
        try:
            return fx.parse_number_input(value)[0]
        except fx.FormulaError:
            return None
    return None


def _same(a: float | None, b: float | None) -> bool:
    if a is None or b is None:
        return a is None and b is None
    return math.isclose(a, b, rel_tol=1e-9, abs_tol=1e-6)


def import_workbook(data: bytes | BinaryIO, columns: list[Column]) -> ImportResult:
    """Read a NetWorth workbook. Unknown header columns become new columns (number columns as input, else text)."""
    raw = data if isinstance(data, bytes) else data.read()
    wb_formulas = openpyxl.load_workbook(io.BytesIO(raw))
    wb_values = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
    sheet = SHEET if SHEET in wb_formulas.sheetnames else wb_formulas.sheetnames[0]
    ws_f, ws_v = wb_formulas[sheet], wb_values[sheet]

    headers = [str(c.value).strip() if c.value is not None else "" for c in ws_f[1]]
    if DATE_COLUMN not in headers:
        raise ValueError(f"Spalte {DATE_COLUMN!r} fehlt in der ersten Zeile von {sheet!r}")
    date_index = headers.index(DATE_COLUMN)

    rows = []
    for formulas, values in zip(ws_f.iter_rows(min_row=2), ws_v.iter_rows(min_row=2)):
        date = _to_date(values[date_index].value)
        if date is not None:
            rows.append((date, {h: (f.value, v.value) for h, f, v in zip(headers, formulas, values) if h}))

    by_name = {c.name: c for c in columns}
    new_columns, ignored = [], []
    for header in headers:
        if not header or header == DATE_COLUMN or header in by_name:
            continue
        cells = [r[1].get(header, (None, None))[1] for r in rows]
        if all(v is None for v in cells):
            ignored.append(header)
            continue
        numeric = all(v is None or isinstance(v, (int, float)) for v in cells)
        new_columns.append(Column(name=header, kind="input" if numeric else "text", width=12.0))
    all_columns = columns + [Column(**{**c.__dict__, "id": -i}) for i, c in enumerate(new_columns, start=1)]

    entries = []
    for date, cells in sorted(rows, key=lambda r: r[0]):  # stable: same-day rows keep their order
        entry = Entry(date=date)
        for col in all_columns:
            if col.name not in cells or col.kind == "formula":
                continue
            formula_text, value = cells[col.name]
            if col.kind == "text":
                if value is not None and str(value).strip():
                    entry.cells[col.id] = Cell(text=str(value).strip())
                continue
            expr = ""
            if isinstance(formula_text, str) and formula_text.startswith("="):
                body = formula_text[1:].strip()
                if fx.is_number_expression(body):
                    expr = body
                    if value is None:  # never calculated (e.g. written by a script)
                        value = fx.parse_number_input(body)[0]
            number = _to_number(value)
            if number is not None:
                entry.cells[col.id] = Cell(value=number, expr=expr)
        entries.append((entry, cells))

    # compute in date order; a formula cell differing from the configured formula is kept as a fixed value
    calc = Calculator(all_columns)
    fixed = 0
    for entry, cells in entries:
        def decide(col: Column, computed, cells=cells, entry=entry):
            nonlocal fixed
            if col.name not in cells:
                return computed, False
            excel_value = _to_number(cells[col.name][1])
            if _same(excel_value, computed):
                return computed, False
            entry.cells[col.id] = Cell(value=excel_value, override=True)
            fixed += 1
            return excel_value, True
        calc.compute(entry, decide)

    return ImportResult(entries=[e for e, _ in entries], income=_import_income(wb_values),
                        new_columns=new_columns, ignored=ignored, fixed_values=fixed)


def _import_income(wb) -> list[Income] | None:
    if INCOME_SHEET not in wb.sheetnames:
        return None
    ws = wb[INCOME_SHEET]
    headers = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    if "Datum" not in headers or "Nettogehalt" not in headers:
        return None
    d, a = headers.index("Datum"), headers.index("Nettogehalt")
    items = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        date, amount = _to_date(row[d]), _to_number(row[a])
        if date is not None and amount is not None:
            items.append(Income(date, amount))
    return sorted(items, key=lambda i: i.date)
