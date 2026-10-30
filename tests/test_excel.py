import io

import openpyxl
import pandas as pd
import pytest

from networth.columns import Column, Layout
from networth.dashboard import render
from networth.excel import export_workbook, import_workbook
from networth.sheet import compute, to_frame


def _import(store, path):
    result = import_workbook(path.read_bytes(), store.columns())
    store.replace_all(result.new_columns, result.entries, result.income)
    return result


def _export(store) -> bytes:
    columns, entries = store.columns(), store.entries()
    return export_workbook(columns, entries, compute(columns, entries), store.income())


def test_import_test_data(store, example_xlsx):
    result = _import(store, example_xlsx)
    original = pd.read_excel(example_xlsx, sheet_name="Net Worth")
    assert len(result.entries) == len(original)
    assert result.new_columns == []
    assert len(result.income) == len(pd.read_excel(example_xlsx, sheet_name="Income"))

    columns = store.columns()
    df = to_frame(columns, compute(columns, store.entries()))
    pd.testing.assert_series_equal(df["Summe"], original["Summe"], check_names=False, atol=0.01)
    assert df["Kommentar"].notna().sum() == original["Kommentar"].notna().sum()


def test_export_has_original_layout_and_cached_values(store, example_xlsx):
    _import(store, example_xlsx)
    content = _export(store)

    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert wb.sheetnames == ["Net Worth", "Income"]
    ws = wb["Net Worth"]
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["Datum", "Girokonto", "Kreditkarte", "Sparkonto"] and headers[-1] == "Kommentar"
    summe = headers.index("Summe") + 1
    assert ws.cell(2, summe).value == "=B2+C2+D2+E2+F2+G2+I2+J2+K2+L2+M2"  # all accounts and the tax
    assert ws.cell(2, 2).number_format == '"€"\\ #,##0.00'

    # pandas (and the old visualize_net_worth.py) read the cached formula results
    exported = pd.read_excel(io.BytesIO(content), sheet_name="Net Worth")
    original = pd.read_excel(content and io.BytesIO(content), sheet_name="Income")
    assert exported["Summe"].notna().all()
    assert len(original) > 0
    source = pd.read_excel(example_xlsx, sheet_name="Net Worth")
    pd.testing.assert_series_equal(exported["Summe"], source["Summe"], atol=0.01)


def test_roundtrip_keeps_expressions_and_fixed_values(store, example_xlsx, tmp_path):
    # an input typed as a sum and an old row whose formula column was a hard value
    wb = openpyxl.load_workbook(example_xlsx)
    ws = wb["Net Worth"]
    headers = [c.value for c in ws[1]]
    ws.cell(3, headers.index("Sparkonto") + 1).value = "=5000+12.5"
    ws.cell(4, headers.index("Summe") + 1).value = 123.45
    path = tmp_path / "modified.xlsx"
    wb.save(path)

    result = _import(store, path)
    assert result.fixed_values >= 1
    sparkonto = next(c for c in store.columns() if c.name == "Sparkonto")
    assert store.entries()[1].cells[sparkonto.id].expr == "5000+12.5"
    assert store.entries()[1].cells[sparkonto.id].value == 5012.5

    exported = tmp_path / "exported.xlsx"
    exported.write_bytes(_export(store))
    ws = openpyxl.load_workbook(exported)["Net Worth"]
    headers = [c.value for c in ws[1]]
    assert ws.cell(3, headers.index("Sparkonto") + 1).value == "=5000+12.5"
    assert ws.cell(4, headers.index("Summe") + 1).value == 123.45

    # importing the export again changes nothing
    before = [r.values for r in compute(store.columns(), store.entries())]
    again = _import(store, exported)
    after = [r.values for r in compute(store.columns(), store.entries())]
    assert again.fixed_values == result.fixed_values
    assert after == before


def test_unknown_columns_are_added(store, example_xlsx, tmp_path):
    df = pd.read_excel(example_xlsx, sheet_name="Net Worth")
    df["Krypto"] = 100.0
    path = tmp_path / "with_new.xlsx"
    df.to_excel(path, sheet_name="Net Worth", index=False)
    result = _import(store, path)
    assert [c.name for c in result.new_columns] == ["Krypto"]
    assert result.income is None
    krypto = next(c for c in store.columns() if c.name == "Krypto")
    assert krypto.spec == "value"  # not part of the net worth until the user makes it an account
    assert all(e.cells[krypto.id].value == 100.0 for e in store.entries())
    summe = next(c for c in store.columns() if c.spec == "total")
    assert not any(e.cells.get(summe.id) for e in store.entries())  # the total is unchanged: no fixed values


def test_new_formula_column_and_dashboard(store, example_xlsx):
    _import(store, example_xlsx)
    store.save_column(Column(name="Krypto", asset_group="Krypto"))
    col = store.save_column(Column(name="Summe ohne Vorsorge", kind="formula",
                                   formula="{Summe}-SUM({Bausparer},{Bundespensionskasse},{Allianz Vorsorgekasse})"))
    with pytest.raises(ValueError, match="Bereiche"):  # ranges depended on the column order
        store.save_column(Column(name="Bereich", kind="formula", formula="SUM({Bausparer}:{Bundesschatz})"))
    columns, entries = store.columns(), store.entries()
    rows = compute(columns, entries)
    assert rows[-1].values[col.name] == pytest.approx(
        rows[-1].values["Summe"] - sum(rows[-1].values[n] for n in ("Bausparer", "Bundespensionskasse",
                                                                     "Allianz Vorsorgekasse")))
    page = render(to_frame(columns, rows), Layout.from_columns(columns, store.group_order()), None, 1e6,
                  plotly_src="plotly.js")
    assert "Vermögensentwicklung" in page and "Krypto" in page


def test_rename_updates_formulas(store):
    giro = next(c for c in store.columns() if c.name == "Girokonto")
    giro.name = "Giro"
    store.save_column(giro)
    by_name = {c.name: c for c in store.columns()}
    assert by_name["Summe"].formula.startswith("{Giro}+{Kreditkarte}+")
    assert "{Giro}" in by_name["Summe Geld"].formula
    assert by_name["Kreditkarte"].net_into == "Giro"


def test_delete_referenced_column_is_refused(store):
    by_name = {c.name: c for c in store.columns()}
    with pytest.raises(ValueError, match="verwendet von: Kreditkarte"):  # the debts booked against it
        store.delete_column(by_name["Girokonto"].id)
    with pytest.raises(ValueError, match="fest dazu"):
        store.delete_column(by_name["Kest"].id)
